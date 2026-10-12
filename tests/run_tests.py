"""End-to-end test of every tool in a real Chromium browser (run by .github/workflows/test.yml).
Serves the site on http://localhost:8000, feeds each tool a sample file, and writes tests/report.json.
"""
import asyncio, json, os, subprocess, sys, time, zipfile
from playwright.async_api import async_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
S = os.path.join(HERE, "samples")
BASE = "http://localhost:8000/"


def make_samples():
    os.makedirs(S, exist_ok=True)
    from PIL import Image, ImageDraw, ImageFont
    im = Image.new("RGB", (1600, 1000), "white"); d = ImageDraw.Draw(im)
    for i in range(0, 1600, 40):
        d.rectangle((i, 600, i + 20, 1000), fill=(i % 255, 120, 200))
    try:
        f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 90)
    except Exception:
        f = ImageFont.load_default()
    d.text((80, 120), "Tools for World", fill="black", font=f)
    d.text((80, 300), "Free online tools", fill="black", font=f)
    d.ellipse((1200, 80, 1500, 380), fill=(220, 60, 60))
    im.save(f"{S}/photo.jpg", quality=92); im.save(f"{S}/photo.png"); im.save(f"{S}/photo.webp")
    try:
        import pillow_heif
        pillow_heif.register_heif_opener(); im.save(f"{S}/photo.heic")
    except Exception as e:
        print("heic sample failed", e)
    from fpdf import FPDF
    for name, pages in (("doc.pdf", 3), ("doc2.pdf", 2)):
        pdf = FPDF(); pdf.set_font("Helvetica", size=14)
        for p in range(pages):
            pdf.add_page(); pdf.set_font("Helvetica", "B", 24); pdf.cell(0, 14, f"Heading page {p + 1}", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", size=12)
            pdf.multi_cell(0, 7, "This is a sample paragraph for testing the PDF tools. " * 12)
            pdf.image(f"{S}/photo.jpg", w=160)
        pdf.output(f"{S}/{name}")
    # scan-like PDF (big noisy photos) for the compressor
    import random
    rnd = random.Random(1); big = Image.new("RGB", (2480, 3508), (245, 243, 238)); bd = ImageDraw.Draw(big)
    for _ in range(9000):
        x, y = rnd.randrange(2480), rnd.randrange(3508); c = rnd.randrange(60, 230)
        bd.rectangle((x, y, x + rnd.randrange(3, 40), y + rnd.randrange(3, 12)), fill=(c, c, c))
    big.save(f"{S}/scan.jpg", quality=95)
    pdf = FPDF()
    for _ in range(3):
        pdf.add_page(); pdf.image(f"{S}/scan.jpg", x=0, y=0, w=210, h=297)
    pdf.output(f"{S}/scan.pdf")
    import docx
    doc = docx.Document(); doc.add_heading("Sample Word document", 1)
    doc.add_paragraph("Hello from Tools for World. This paragraph tests Word to PDF conversion. " * 5)
    t = doc.add_table(rows=2, cols=2); t.cell(0, 0).text = "A"; t.cell(1, 1).text = "D"
    doc.save(f"{S}/doc.docx")
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=25", "-f", "lavfi",
                    "-i", "sine=frequency=440:sample_rate=44100", "-t", "6", "-c:v", "libx264", "-pix_fmt", "yuv420p",
                    "-c:a", "aac", "-shortest", f"{S}/video.mp4"], check=True)


FILE_TESTS = [  # slug, sample, timeout s, extra setup
    ("compress-image", "photo.jpg", 30, None), ("compress-image-to-50kb", "photo.jpg", 60, None), ("resize-image-to-1080x1080", "photo.jpg", 30, None), ("resize-image", "photo.jpg", 30, ("#rw", "800")),
    ("png-to-jpg", "photo.png", 30, None), ("jpg-to-png", "photo.jpg", 30, None), ("webp-to-jpg", "photo.webp", 30, None),
    ("heic-to-jpg", "photo.heic", 90, None), ("passport-photo-maker", "photo.jpg", 30, None),
    ("jpg-to-pdf", "photo.jpg", 30, None), ("merge-pdf", ["doc.pdf", "doc2.pdf"], 60, None),
    ("split-pdf", "doc.pdf", 60, ("#sr", "1-2")), ("pdf-to-jpg", "doc.pdf", 90, None), ("compress-pdf", "scan.pdf", 180, None),
    ("pdf-to-word", "doc.pdf", 90, None), ("word-to-pdf", "doc.docx", 120, None), ("image-to-text", "photo.png", 240, None),
    ("mp4-to-mp3", "video.mp4", 120, None), ("remove-background", "photo.jpg", 420, None),
]


async def main():
    make_samples()
    srv = subprocess.Popen([sys.executable, "-m", "http.server", "8000"], cwd=os.path.dirname(HERE), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    time.sleep(2)
    report = {}
    async with async_playwright() as p:
        b = await p.chromium.launch()
        ctx = await b.new_context(accept_downloads=True)
        for slug, sample, tmo, setup in FILE_TESTS:
            pg = await ctx.new_page(); errs = []
            pg.on("pageerror", lambda e, errs=errs: errs.append(str(e)))
            pg.on("console", lambda m, errs=errs: errs.append("console: " + m.text[:200]) if m.type in ("warning", "error") else None)
            t0 = time.time(); r = {"ok": False}
            try:
                await pg.goto(BASE + slug + "/", wait_until="networkidle", timeout=60000)
                files = [f"{S}/{x}" for x in (sample if isinstance(sample, list) else [sample])]
                await pg.set_input_files("#fin", files); await pg.wait_for_timeout(600)
                if setup:
                    await pg.fill(setup[0], setup[1])
                await pg.click("#go")
                await pg.wait_for_function("() => { const s = document.querySelector('#status'); return s && (/Done|smaller/.test(s.textContent) || s.classList.contains('err')); }", timeout=tmo * 1000)
                st = await pg.inner_text("#status"); out = await pg.inner_text("#out")
                r.update(status=st, out=out.replace("\n", " | ")[:300], ok=("err" not in (await pg.get_attribute("#status", "class") or "")) and "Download" in out)
                link = pg.locator("#out a.save").first
                if await link.count():
                    async with pg.expect_download() as dl:
                        await link.click()
                    path = await (await dl.value).path(); r["download_bytes"] = os.path.getsize(path)
                    if slug == "pdf-to-word":
                        with zipfile.ZipFile(path) as z:
                            xml = z.read("word/document.xml").decode(); r["docx_has_text"] = "sample paragraph" in xml
                            r["ok"] = r["ok"] and r["docx_has_text"]
                if slug == "image-to-text":
                    txt = await pg.input_value("#ocrout"); r["ocr_text"] = txt[:120]; r["ok"] = r["ok"] and "Tools" in txt
            except Exception as e:
                r["error"] = str(e)[:300]
                try: r["status"] = await pg.inner_text("#status")
                except Exception: pass
            r["page_errors"] = errs[:5]; r["seconds"] = round(time.time() - t0, 1)
            report[slug] = r; print(slug, "OK" if r["ok"] else "FAIL", r.get("status"), r.get("error", ""))
            await pg.close()
        # calculators / utilities
        pg = await ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
        checks = {}
        await pg.goto(BASE + "qr-code-generator/", wait_until="networkidle")
        await pg.wait_for_timeout(800)
        checks["qr-code-generator"] = {"info": await pg.inner_text("#qinfo"), "dark_pixels": await pg.evaluate(
            "(()=>{const c=document.querySelector('#qrc');const d=c.getContext('2d').getImageData(0,0,c.width,c.height).data;let n=0;for(let i=0;i<d.length;i+=4)if(d[i]<50)n++;return n})()")}
        checks["qr-code-generator"]["ok"] = checks["qr-code-generator"]["dark_pixels"] > 1000
        for slug, sel in (("word-counter", "#out-kv"), ("age-calculator", "#out-total"), ("password-generator", "#out-strength")):
            await pg.goto(BASE + slug + "/", wait_until="networkidle"); await pg.wait_for_timeout(300)
            txt = await pg.inner_text(sel); checks[slug] = {"text": txt.replace("\n", " ")[:120], "ok": bool(txt.strip())}
        await pg.goto(BASE + "percentage-calculator/"); await pg.wait_for_timeout(300)
        r1 = await pg.inner_text("#r1"); checks["percentage-calculator"] = {"r1": r1, "ok": r1 == "30"}
        for k, v in checks.items():
            v["page_errors"] = errs[:3]; report[k] = v; print(k, "OK" if v["ok"] else "FAIL", v)
        await b.close()
    srv.terminate()
    report["_summary"] = {"passed": sum(1 for k, v in report.items() if not k.startswith("_") and v.get("ok")), "total": len(report) - 0,
                          "time": time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime())}
    json.dump(report, open(os.path.join(HERE, "report.json"), "w"), indent=1)
    print(report["_summary"])


asyncio.run(main())
