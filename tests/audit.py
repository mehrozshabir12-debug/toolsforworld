"""Deep audit of every file tool: real sample files, edge cases, output validation,
network capture (proves files are not uploaded), mobile layout and console errors.
Run by .github/workflows/audit.yml; writes tests/audit_report.json.
Needs: playwright, pillow, pillow-heif, fpdf2, python-docx, pypdf; apt: ffmpeg, qpdf, poppler-utils.
"""
import asyncio, json, os, random, re, subprocess, sys, time, zipfile, io
from urllib.parse import urlparse
from playwright.async_api import async_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
S = os.path.join(HERE, "audit_samples")
OUT = os.path.join(HERE, "audit_out")
BASE = "http://localhost:8001/"
RAW_JS_ERR = re.compile(r"undefined|null|Cannot read|is not a function|is not defined|NaN|\[object|Failed to execute|Overload|TypeError|RangeError", re.I)


# ------------------------------------------------------------------ samples
def make_samples():
    os.makedirs(S, exist_ok=True); os.makedirs(OUT, exist_ok=True)
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 90)
    except Exception:
        font = ImageFont.load_default()

    def photo(w, h, path, **kw):
        im = Image.new("RGB", (w, h), "white"); d = ImageDraw.Draw(im)
        for i in range(0, w, max(1, w // 40)):
            d.rectangle((i, int(h * .6), i + max(1, w // 80), h), fill=(i % 255, 120, 200))
        d.text((w // 20, h // 8), "Tools for World", fill="black", font=font)
        d.text((w // 20, h // 3), "Free online tools", fill="black", font=font)
        d.ellipse((int(w * .75), h // 12, int(w * .93), int(h * .38)), fill=(220, 60, 60))
        im.save(path, **kw); return im

    im = photo(1600, 1000, f"{S}/photo.jpg", quality=92)
    im.save(f"{S}/photo.png"); im.save(f"{S}/photo.webp")
    import shutil; shutil.copy(f"{S}/photo.jpg", f"{S}/photo.jfif"); shutil.copy(f"{S}/photo.jpg", f"{S}/photo.jpeg")
    open(f"{S}/notes.txt", "w", encoding="utf-8").write("Tools for World sample paragraph.\n\n" + "A long line that keeps going so it has to wrap onto the next line of the PDF page. " * 6 + "\n" + "\n".join(f"Line {i}" for i in range(1, 121)))
    open(f"{S}/urdu.txt", "w", encoding="utf-8").write("یہ ایک نمونہ ہے\nTools for World\n")
    # 48 MP photo like a modern phone camera (noisy so it compresses badly)
    big = photo(8000, 6000, f"{S}/tmp.png")
    rnd = random.Random(2); d = ImageDraw.Draw(big)
    for _ in range(40000):
        x, y = rnd.randrange(8000), rnd.randrange(6000); d.point((x, y), fill=(rnd.randrange(256),) * 3)
    big.save(f"{S}/huge-48mp.jpg", quality=90); os.remove(f"{S}/tmp.png")
    # transparent PNG
    a = Image.new("RGBA", (800, 600), (0, 0, 0, 0)); ImageDraw.Draw(a).ellipse((100, 100, 700, 500), fill=(30, 140, 90, 255))
    a.save(f"{S}/alpha.png"); a.save(f"{S}/alpha.webp", lossless=True)
    # portrait person-like image for passport / bg removal
    p = Image.new("RGB", (900, 1200), (200, 210, 225)); d = ImageDraw.Draw(p)
    d.ellipse((300, 200, 600, 560), fill=(225, 185, 150)); d.rectangle((220, 560, 680, 1200), fill=(40, 60, 120))
    p.save(f"{S}/portrait.jpg", quality=92)
    open(f"{S}/corrupt.jpg", "wb").write(b"\xff\xd8\xff\xe0" + os.urandom(5000))
    open(f"{S}/empty.jpg", "wb").write(b"")
    try:
        import pillow_heif
        pillow_heif.register_heif_opener(); im.save(f"{S}/photo.heic")
    except Exception as e:
        print("heic sample failed", e)

    try:
        make_docs()
    except ImportError as e:
        print("skipping PDF/Word samples:", e)
    make_media()


def make_docs():
    from PIL import Image, ImageDraw
    from fpdf import FPDF
    pdf = FPDF(); pdf.set_font("Helvetica", size=14)
    for p_ in range(3):
        pdf.add_page(); pdf.set_font("Helvetica", "B", 24); pdf.cell(0, 14, f"Heading page {p_ + 1}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=12); pdf.multi_cell(0, 7, "This is a sample paragraph for testing the PDF tools. " * 12)
        pdf.image(f"{S}/photo.jpg", w=160)
    pdf.output(f"{S}/doc.pdf")
    pdf = FPDF(); pdf.set_font("Helvetica", size=12)
    for p_ in range(60):
        pdf.add_page(); pdf.multi_cell(0, 7, f"Page {p_ + 1}. " + "Long document body text for stress testing. " * 30)
    pdf.output(f"{S}/long-60p.pdf")
    # scan-only PDF (no text layer)
    big2 = Image.new("RGB", (2480, 3508), (245, 243, 238)); bd = ImageDraw.Draw(big2); r2 = random.Random(1)
    for _ in range(9000):
        x, y = r2.randrange(2480), r2.randrange(3508); c = r2.randrange(60, 230)
        bd.rectangle((x, y, x + r2.randrange(3, 40), y + r2.randrange(3, 12)), fill=(c, c, c))
    big2.save(f"{S}/scan.jpg", quality=95)
    pdf = FPDF()
    for _ in range(3):
        pdf.add_page(); pdf.image(f"{S}/scan.jpg", x=0, y=0, w=210, h=297)
    pdf.output(f"{S}/scan.pdf")
    subprocess.run(["qpdf", "--encrypt", "user", "owner", "256", "--", f"{S}/doc.pdf", f"{S}/encrypted.pdf"], check=False)
    subprocess.run(["qpdf", "--rotate=+90:1", f"{S}/doc.pdf", f"{S}/rotated.pdf"], check=False)
    data = open(f"{S}/doc.pdf", "rb").read(); open(f"{S}/corrupt.pdf", "wb").write(data[: len(data) // 3])
    open(f"{S}/fake.pdf", "wb").write(b"this is not a pdf at all")

    import docx
    from docx.shared import Inches
    doc = docx.Document(); doc.add_heading("Sample Word document", 1)
    doc.add_paragraph("Hello from Tools for World. This paragraph tests Word to PDF conversion. " * 5)
    t = doc.add_table(rows=2, cols=2); t.cell(0, 0).text = "A"; t.cell(1, 1).text = "D"
    doc.save(f"{S}/doc.docx")
    doc = docx.Document()
    for i in range(45):
        doc.add_heading(f"Chapter {i + 1}", 1)
        for _ in range(4):
            doc.add_paragraph(f"Chapter {i + 1} text. " + "A long report paragraph that fills the page with words. " * 14)
        if i % 10 == 0:
            doc.add_picture(f"{S}/photo.jpg", width=Inches(4))
    doc.add_paragraph("FINAL-LINE-MARKER the end of the long document.")
    doc.save(f"{S}/long-45ch.docx")
    open(f"{S}/old.doc", "wb").write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + os.urandom(3000))
    zb = io.BytesIO(); zipfile.ZipFile(zb, "w").writestr("hello.txt", "not a docx"); open(f"{S}/corrupt.docx", "wb").write(zb.getvalue())



def make_media():
    ff = ["ffmpeg", "-y", "-loglevel", "error"]
    subprocess.run(ff + ["-f", "lavfi", "-i", "testsrc=size=640x360:rate=25", "-f", "lavfi", "-i", "sine=frequency=440:sample_rate=48000",
                         "-t", "6", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-ac", "2", "-shortest", f"{S}/video.mp4"], check=True)
    subprocess.run(ff + ["-f", "lavfi", "-i", "testsrc=size=320x240:rate=25", "-t", "4", "-c:v", "libx264", "-pix_fmt", "yuv420p", f"{S}/noaudio.mp4"], check=True)
    subprocess.run(ff + ["-f", "lavfi", "-i", "testsrc=size=160x120:rate=5", "-f", "lavfi", "-i", "sine=frequency=330:sample_rate=44100",
                         "-t", "600", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", f"{S}/long-10min.mp4"], check=True)
    subprocess.run(ff + ["-f", "lavfi", "-i", "sine=frequency=500:sample_rate=44100", "-t", "5", "-c:a", "aac", f"{S}/audio.m4a"], check=True)
    subprocess.run(ff + ["-i", f"{S}/video.mp4", "-c", "copy", f"{S}/video.mov"], check=True)   # iPhone-style MOV
    subprocess.run(ff + ["-i", f"{S}/video.mp4", "-c:v", "libvpx", "-b:v", "300k", "-c:a", "libopus", f"{S}/video.webm"], check=True)   # screen-recorder WebM


# ------------------------------------------------------------------ output validators
def v_image(path, w=None, h=None, max_bytes=None, alpha=False, fmt=None):
    from PIL import Image
    im = Image.open(path); im.load(); r = {"w": im.width, "h": im.height, "fmt": im.format, "mode": im.mode}
    ok = True
    if w and im.width != w: ok = False
    if h and im.height != h: ok = False
    if max_bytes and os.path.getsize(path) > max_bytes: ok = False
    if fmt and im.format != fmt: ok = False
    if alpha:
        a = im.convert("RGBA").getchannel("A"); lo, hi = a.getextrema(); r["alpha_range"] = [lo, hi]; ok = ok and lo < 20 and hi > 235
    return ok, r


def v_pdf(path, pages=None, check_last_page_not_blank=False, text=None, min_pages=None, max_bytes=None):
    from pypdf import PdfReader
    rd = PdfReader(path); n = len(rd.pages); r = {"pages": n, "bytes": os.path.getsize(path)}; ok = True
    if max_bytes and r["bytes"] > max_bytes: ok = False
    if pages is not None and n != pages: ok = False
    if min_pages is not None and n < min_pages: ok = False
    if text is not None:
        t = "".join((pg.extract_text() or "") for pg in rd.pages); r["has_text"] = text in t; ok = ok and r["has_text"]
    if check_last_page_not_blank:
        from PIL import Image
        blanks = []
        for i in sorted({1, max(1, n // 2), n}):
            subprocess.run(["pdftoppm", "-f", str(i), "-l", str(i), "-r", "30", "-png", path, f"{OUT}/chk"], check=True)
            png = [f for f in os.listdir(OUT) if f.startswith("chk")][0]
            im = Image.open(f"{OUT}/{png}").convert("L"); hist = im.histogram(); dark = sum(hist[:200]) / (im.width * im.height)
            os.remove(f"{OUT}/{png}"); blanks.append([i, round(dark, 4)])
        r["dark_share_by_page"] = blanks; ok = ok and all(d > 0.002 for _, d in blanks)
    return ok, r


def v_docx(path, text):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode()
    ok = text in xml; return ok, {"has_text": ok, "xml_kb": len(xml) // 1024}


def v_mp3(path, dur):
    p = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration:stream=codec_name,sample_rate,channels", "-of", "json", path], capture_output=True, text=True)
    j = json.loads(p.stdout or "{}"); d = float(j.get("format", {}).get("duration", 0)); st = (j.get("streams") or [{}])[0]
    ok = st.get("codec_name") == "mp3" and abs(d - dur) < max(1.0, dur * 0.03)
    return ok, {"duration": round(d, 2), "codec": st.get("codec_name"), "rate": st.get("sample_rate"), "ch": st.get("channels")}


def v_text(path, needle):
    t = open(path, encoding="utf-8", errors="replace").read(); return needle in t, {"text": t[:80]}


# ------------------------------------------------------------------ cases
# (id, slug, files, expect, setup, validators[list of (output_index, fn, kwargs)], timeout_s)
CASES = [
    # images
    ("compress-basic", "compress-image", ["photo.jpg"], "ok", {}, [(0, "image", {"fmt": "JPEG"})], 60),
    ("compress-to-50kb", "compress-image-to-50kb", ["photo.jpg"], "ok", {}, [(0, "image", {"max_bytes": 50 * 1024})], 60),
    ("compress-to-20kb-huge", "compress-image-to-20kb", ["huge-48mp.jpg"], "ok", {}, [(0, "image", {"max_bytes": 20 * 1024})], 240),
    ("compress-huge-48mp", "compress-image", ["huge-48mp.jpg"], "ok", {}, [(0, "image", {})], 240),
    ("compress-batch-3", "compress-image", ["photo.jpg", "photo.png", "photo.webp"], "ok", {}, [(2, "image", {})], 90),
    ("compress-alpha-png", "compress-image", ["alpha.png"], "ok", {}, [(0, "image", {})], 60),
    ("compress-corrupt", "compress-image", ["corrupt.jpg"], "err", {}, [], 60),
    ("compress-empty", "compress-image", ["empty.jpg"], "err", {}, [], 60),
    ("compress-heic-wrong-tool", "compress-image", ["photo.heic"], "err_or_reject", {}, [], 60),
    ("compress-pdf-wrong-type", "compress-image", ["doc.pdf"], "reject", {}, [], 30),
    ("resize-exact-1080", "resize-image-to-1080x1920", ["photo.jpg"], "ok", {}, [(0, "image", {"w": 1080, "h": 1920})], 60),
    ("resize-exact-pad", "resize-image-to-1920x1080", ["portrait.jpg"], "ok", {"#rf": "pad"}, [(0, "image", {"w": 1920, "h": 1080})], 60),
    ("resize-width-only", "resize-image", ["photo.jpg"], "ok", {"#rw": "800"}, [(0, "image", {"w": 800, "h": 500})], 60),
    ("resize-percent-400-huge", "resize-image", ["huge-48mp.jpg"], "ok_or_err", {"#rm": "pc", "#rp": "400"}, [], 240),
    ("resize-negative", "resize-image", ["photo.jpg"], "err", {"#rw": "-50"}, [], 30),
    ("resize-empty-size", "resize-image", ["photo.jpg"], "err", {}, [], 30),
    ("png-to-jpg-alpha", "png-to-jpg", ["alpha.png"], "ok", {}, [(0, "image", {"fmt": "JPEG", "w": 800})], 60),
    ("jpg-to-png", "jpg-to-png", ["photo.jpg"], "ok", {}, [(0, "image", {"fmt": "PNG", "w": 1600})], 60),
    ("webp-to-jpg", "webp-to-jpg", ["photo.webp"], "ok", {}, [(0, "image", {"fmt": "JPEG"})], 60),
    ("jfif-to-jpg", "jfif-to-jpg", ["photo.jfif"], "ok", {}, [(0, "image", {"fmt": "JPEG", "w": 1600})], 60),
    ("jpeg-to-jpg", "jpeg-to-jpg", ["photo.jpeg"], "ok", {}, [(0, "image", {"fmt": "JPEG", "w": 1600})], 60),
    ("webp-to-png-alpha", "webp-to-png", ["alpha.webp"], "ok", {}, [(0, "image", {"fmt": "PNG", "w": 800, "alpha": True})], 60),
    ("webp-to-png-wrong-type", "webp-to-png", ["photo.jpg"], "reject", {}, [], 30),
    ("heic-to-jpg", "heic-to-jpg", ["photo.heic"], "ok", {}, [(0, "image", {"fmt": "JPEG", "w": 1600})], 120),
    ("heic-corrupt", "heic-to-jpg", ["corrupt.jpg"], "reject", {}, [], 30),
    ("passport-uk", "uk-passport-photo", ["portrait.jpg"], "ok", {}, [(0, "image", {"w": 413, "h": 531}), (1, "image", {"w": 1800, "h": 1200})], 60),
    ("passport-us", "us-passport-photo", ["portrait.jpg"], "ok", {}, [(0, "image", {"w": 600, "h": 600})], 60),
    ("passport-corrupt", "passport-photo-maker", ["corrupt.jpg"], "err_or_reject", {}, [], 60),
    ("bgremove", "remove-background", ["portrait.jpg"], "ok", {}, [(0, "image", {"alpha": True, "fmt": "PNG"})], 600),
    ("ocr-eng", "image-to-text", ["photo.png"], "ok", {}, [(0, "text", {"needle": "Tools"})], 300),
    ("ocr-corrupt", "image-to-text", ["corrupt.jpg"], "err", {}, [], 300),
    # pdf
    ("jpg-to-pdf-3", "jpg-to-pdf", ["photo.jpg", "photo.png", "photo.webp"], "ok", {}, [(0, "pdf", {"pages": 3})], 60),
    ("jpg-to-pdf-huge", "jpg-to-pdf", ["huge-48mp.jpg"], "ok", {}, [(0, "pdf", {"pages": 1})], 240),
    ("jpg-to-pdf-under-200kb-3", "jpg-to-pdf-under-200kb", ["photo.jpg", "photo.png", "photo.webp"], "ok", {}, [(0, "pdf", {"pages": 3, "max_bytes": 200 * 1024})], 180),
    ("jpg-to-pdf-under-100kb-huge", "jpg-to-pdf-under-100kb", ["huge-48mp.jpg"], "ok", {}, [(0, "pdf", {"pages": 1, "max_bytes": 100 * 1024})], 300),
    ("merge-2", "merge-pdf", ["doc.pdf", "long-60p.pdf"], "ok", {}, [(0, "pdf", {"pages": 63})], 120),
    ("merge-encrypted", "merge-pdf", ["doc.pdf", "encrypted.pdf"], "ok_or_err", {}, [], 60),
    ("merge-corrupt", "merge-pdf", ["doc.pdf", "fake.pdf"], "err", {}, [], 60),
    ("split-range", "split-pdf", ["long-60p.pdf"], "ok", {"#sr": "2-4, 10"}, [(0, "pdf", {"pages": 4})], 60),
    ("split-bad-range", "split-pdf", ["doc.pdf"], "err", {"#sr": "50-60"}, [], 60),
    ("split-encrypted", "split-pdf", ["encrypted.pdf"], "ok_or_err", {"#sr": "1"}, [], 60),
    ("pdf-to-jpg", "pdf-to-jpg", ["doc.pdf"], "ok", {}, [(0, "image", {"fmt": "JPEG"})], 120),
    ("pdf-to-jpg-rotated", "pdf-to-jpg", ["rotated.pdf"], "ok", {}, [(0, "image", {})], 120),
    ("pdf-to-jpg-encrypted", "pdf-to-jpg", ["encrypted.pdf"], "err", {}, [], 60),
    ("pdf-to-png", "pdf-to-png", ["doc.pdf"], "ok", {}, [(0, "image", {"fmt": "PNG"})], 120),
    ("pdf-to-png-transparent", "pdf-to-png", ["doc.pdf"], "ok", {"#dt": True}, [(0, "image", {"fmt": "PNG", "alpha": True})], 120),
    ("pdf-to-png-encrypted", "pdf-to-png", ["encrypted.pdf"], "err", {}, [], 60),
    ("png-to-pdf-2", "png-to-pdf", ["alpha.png", "photo.png"], "ok", {}, [(0, "pdf", {"pages": 2})], 60),
    ("heic-to-pdf-mixed", "heic-to-pdf", ["photo.heic", "photo.jpg"], "ok", {}, [(0, "pdf", {"pages": 2})], 180),
    ("txt-to-pdf", "txt-to-pdf", ["notes.txt"], "ok", {}, [(0, "pdf", {"min_pages": 2, "text": "sample paragraph"})], 60),
    ("txt-to-pdf-urdu", "txt-to-pdf", ["urdu.txt"], "ok", {}, [(0, "pdf", {"pages": 1, "check_last_page_not_blank": True})], 60),
    ("txt-to-pdf-wrong-type", "txt-to-pdf", ["photo.jpg"], "reject", {}, [], 30),
    ("pdf-to-jpg-corrupt", "pdf-to-jpg", ["corrupt.pdf"], "err_or_ok", {}, [], 60),
    ("compress-pdf-scan", "compress-pdf", ["scan.pdf"], "ok", {}, [(0, "pdf", {"pages": 3, "check_last_page_not_blank": True})], 240),
    ("compress-pdf-to-200kb", "compress-pdf-to-200kb", ["scan.pdf"], "ok", {}, [(0, "pdf", {"pages": 3})], 400),
    ("compress-pdf-text-only", "compress-pdf", ["long-60p.pdf"], "ok_or_err", {}, [], 400),
    ("compress-pdf-rotated", "compress-pdf", ["rotated.pdf"], "ok_or_err", {}, [], 240),
    ("compress-pdf-fake", "compress-pdf", ["fake.pdf"], "err", {}, [], 60),
    ("pdf-to-word", "pdf-to-word", ["doc.pdf"], "ok", {}, [(0, "docx", {"text": "sample paragraph"})], 120),
    ("pdf-to-word-60p", "pdf-to-word", ["long-60p.pdf"], "ok", {}, [(0, "docx", {"text": "Page 60."})], 180),
    ("pdf-to-word-scan", "pdf-to-word", ["scan.pdf"], "err", {}, [], 120),
    ("pdf-to-word-encrypted", "pdf-to-word", ["encrypted.pdf"], "err", {}, [], 60),
    ("word-to-pdf", "word-to-pdf", ["doc.docx"], "ok", {}, [(0, "pdf", {"check_last_page_not_blank": True})], 180),
    ("word-to-pdf-long", "word-to-pdf", ["long-45ch.docx"], "ok", {}, [(0, "pdf", {"min_pages": 30, "check_last_page_not_blank": True})], 600),
    ("word-to-pdf-old-doc", "word-to-pdf", ["old.doc"], "reject_or_err", {}, [], 60),
    ("word-to-pdf-corrupt", "word-to-pdf", ["corrupt.docx"], "err", {}, [], 60),
    # media
    ("mp4-to-mp3", "mp4-to-mp3", ["video.mp4"], "ok", {}, [(0, "mp3", {"dur": 6})], 180),
    ("mp4-to-mp3-10min", "mp4-to-mp3", ["long-10min.mp4"], "ok", {}, [(0, "mp3", {"dur": 600})], 600),
    ("m4a-to-mp3", "mp4-to-mp3", ["audio.m4a"], "ok", {}, [(0, "mp3", {"dur": 5})], 120),
    ("mp4-no-audio", "mp4-to-mp3", ["noaudio.mp4"], "err", {}, [], 120),
    ("video-to-mp3-mov", "video-to-mp3", ["video.mov"], "ok", {}, [(0, "mp3", {"dur": 6})], 180),
    ("video-to-mp3-webm", "video-to-mp3", ["video.webm"], "ok", {}, [(0, "mp3", {"dur": 6})], 180),
]
VALIDATORS = {"image": v_image, "pdf": v_pdf, "docx": v_docx, "mp3": v_mp3, "text": v_text}
ANALYTICS = ("cloudflareinsights.com", "googletagmanager.com", "google-analytics.com", "googlesyndication.com", "doubleclick.net", "google.com", "googleadservices.com",
             "gstatic.com", "googleapis.com", "fundingchoicesmessages.google.com", "adtrafficquality.google")


async def run_case(ctx, case, mobile=False):
    cid, slug, files, expect, setup, vals, tmo = case
    pg = await ctx.new_page(); errs = []; net = []
    pg.on("pageerror", lambda e: errs.append("pageerror: " + str(e)[:200]))
    pg.on("console", lambda m: errs.append("console: " + m.text[:160]) if m.type == "error" and "cloudflareinsights" not in m.text and "ERR_FAILED" not in m.text else None)

    def on_req(rq):
        u = urlparse(rq.url)
        if u.scheme in ("data", "blob"): return
        body = rq.post_data_buffer
        net.append({"m": rq.method, "host": u.hostname, "path": u.path[:60], "body": len(body) if body else 0})
    pg.on("request", on_req)
    r = {"slug": slug, "files": files, "expect": expect, "mobile": mobile}
    t0 = time.time()
    try:
        await pg.goto(BASE + slug + "/", wait_until="load", timeout=60000)
        await pg.wait_for_timeout(500)
        mem0 = await pg.evaluate("performance.memory ? performance.memory.usedJSHeapSize : 0")
        await pg.set_input_files("#fin", [f"{S}/{x}" for x in files]); await pg.wait_for_timeout(700)
        st_cls = await pg.get_attribute("#status", "class") or ""
        panel_hidden = await pg.evaluate("document.querySelector('#panel').hidden")
        if panel_hidden or ("err" in st_cls and await pg.evaluate("document.querySelectorAll('#list .file').length") == 0):
            r.update(outcome="reject", status=await pg.inner_text("#status"))
        else:
            for sel, v in setup.items():
                el = pg.locator(sel)
                if await el.evaluate("e => e.tagName") == "SELECT": await el.select_option(v)
                elif await el.get_attribute("type") == "checkbox": await (el.check() if v else el.uncheck())
                else: await el.fill(v)
            await pg.click("#go")
            await pg.wait_for_function("() => !document.querySelector('#go').disabled", timeout=tmo * 1000)
            await pg.wait_for_timeout(400)
            st = await pg.inner_text("#status"); cls = await pg.get_attribute("#status", "class") or ""
            r.update(status=st[:300], outcome="err" if "err" in cls else "ok")
            links = pg.locator("#out a.save"); n = await links.count(); r["outputs"] = n
            r["out_text"] = (await pg.inner_text("#out")).replace("\n", " | ")[:300]
            saved = []
            for i in range(min(n, 3)):
                async with pg.expect_download(timeout=60000) as dl:
                    await links.nth(i).click()
                d = await dl.value; p = os.path.join(OUT, f"{cid}-{i}-{d.suggested_filename}"); await d.save_as(p); saved.append(p)
            r["saved"] = [os.path.basename(p) + f" ({os.path.getsize(p)} B)" for p in saved]
            if r["outcome"] == "err" and RAW_JS_ERR.search(st): r["raw_error_message"] = True
            if r["outcome"] == "ok":
                vres = []
                for idx, kind, kw in vals:
                    if idx >= len(saved): vres.append({"ok": False, "why": f"output {idx} missing"}); continue
                    try:
                        ok, info = VALIDATORS[kind](saved[idx], **kw); vres.append({"ok": ok, **info})
                    except Exception as e:
                        vres.append({"ok": False, "why": f"validator error: {e}"[:200]})
                r["validation"] = vres
        mem1 = await pg.evaluate("performance.memory ? performance.memory.usedJSHeapSize : 0")
        r["heap_mb"] = round(mem1 / 1048576, 1); r["heap_growth_mb"] = round((mem1 - mem0) / 1048576, 1)
    except Exception as e:
        r["outcome"] = "exception"; r["error"] = str(e)[:300]
        try: r["status"] = await pg.inner_text("#status")
        except Exception: pass
    r["seconds"] = round(time.time() - t0, 1)
    r["page_errors"] = errs[:6]
    big_uploads = [n for n in net if n["body"] > 4000]
    r["hosts"] = sorted({n["host"] for n in net if n["host"] and n["host"] not in ("localhost",)})
    r["uploads_over_4kb"] = big_uploads
    r["non_get_requests"] = [n for n in net if n["m"] not in ("GET", "HEAD", "OPTIONS") and not any(n["host"].endswith(a) for a in ANALYTICS)]
    # verdict
    out = r.get("outcome"); allowed = expect.split("_or_")
    ok = out in allowed
    if out == "ok" and "validation" in r: ok = ok and all(v.get("ok") for v in r["validation"])
    if r.get("raw_error_message"): ok = False
    if any(e.startswith("pageerror") for e in errs): ok = False
    if big_uploads: ok = False; r["PRIVACY"] = "file-sized request body sent!"
    r["pass"] = ok
    await pg.close()
    return r


async def layout_audit(b, urls):
    """Mobile (360 px) and desktop: horizontal overflow, missing alt, small tap targets, h1 count."""
    res = {}
    for name, vp, mob in (("mobile", {"width": 360, "height": 740}, True), ("desktop", {"width": 1366, "height": 800}, False)):
        ctx = await b.new_context(viewport=vp, is_mobile=mob, has_touch=mob, device_scale_factor=2 if mob else 1)
        pg = await ctx.new_page(); errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)[:150]))
        bad = []
        for u in urls:
            await pg.goto(u, wait_until="load"); await pg.wait_for_timeout(150)
            info = await pg.evaluate("""() => {
              const pn = document.querySelector("#panel"); if (pn) pn.hidden = false;
              const de = document.documentElement, over = de.scrollWidth - de.clientWidth;
              const wide = [...document.querySelectorAll('body *')].filter(e => { const r = e.getBoundingClientRect(); return r.right > de.clientWidth + 1 && getComputedStyle(e).position !== 'fixed' && r.width > 0; }).slice(0, 3).map(e => e.tagName + '.' + e.className + '#' + e.id);
              const small = [...document.querySelectorAll('a, button, input, select, summary')].filter(e => { const r = e.getBoundingClientRect(); return r.width > 0 && r.height > 0 && (r.height < 24 || r.width < 24) && e.closest('main,header,footer'); }).length;
              return {over, wide, small, h1: document.querySelectorAll('h1').length, imgNoAlt: document.querySelectorAll('img:not([alt])').length};
            }""")
            if info["over"] > 0 or info["h1"] != 1 or info["imgNoAlt"]:
                bad.append({"url": u.replace(BASE, "/"), **info})
            res.setdefault(name + "_small_tap_targets", {})[u.replace(BASE, "/")] = info["small"]
        res[name] = {"pages": len(urls), "problems": bad, "page_errors": errs[:5]}
        await ctx.close()
    # keep only the worst tap-target pages
    for k in [k for k in res if k.endswith("_small_tap_targets")]:
        res[k] = dict(sorted(res[k].items(), key=lambda kv: -kv[1])[:8])
    return res


async def main():
    only = sys.argv[1:]
    layout_only = only == ["layout"]
    if layout_only: only = ["-none-"]
    make_samples()
    srv = subprocess.Popen([sys.executable, "-m", "http.server", "8001"], cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    time.sleep(2)
    report = {"time": time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime()), "cases": {}}
    async with async_playwright() as p:
        b = await p.chromium.launch(args=["--enable-precise-memory-info"])
        ctx = await b.new_context(accept_downloads=True, viewport={"width": 1280, "height": 900})
        for c in CASES:
            if only and c[0] not in only: continue
            r = await run_case(ctx, c); report["cases"][c[0]] = r
            print(("PASS " if r["pass"] else "FAIL ") + c[0], r.get("outcome"), r.get("seconds"), (r.get("status") or "")[:90], r.get("error", ""), flush=True)
        await ctx.close()
        # a few tools again on an emulated phone
        mctx = await b.new_context(accept_downloads=True, **p.devices["Pixel 7"])
        for c in CASES:
            if c[0] in ("compress-basic", "pdf-to-jpg", "passport-uk", "word-to-pdf", "pdf-to-word", "mp4-to-mp3") and (not only or c[0] in only):
                r = await run_case(mctx, c, mobile=True); report["cases"]["mobile:" + c[0]] = r
                print(("PASS " if r["pass"] else "FAIL ") + "mobile:" + c[0], r.get("outcome"), r.get("seconds"), flush=True)
        await mctx.close()
        if not only or layout_only:
            urls = re.findall(r"<loc>https://toolsforworld.com/(.*?)</loc>", open(os.path.join(ROOT, "sitemap.xml")).read())
            report["layout"] = await layout_audit(b, [BASE + u for u in urls])
        await b.close()
    srv.terminate()
    cs = report["cases"]
    report["summary"] = {"passed": sum(1 for v in cs.values() if v["pass"]), "total": len(cs),
                         "failed": [k for k, v in cs.items() if not v["pass"]],
                         "privacy_violations": [k for k, v in cs.items() if v.get("PRIVACY")],
                         "third_party_hosts": sorted({h for v in cs.values() for h in v.get("hosts", [])})}
    json.dump(report, open(os.path.join(HERE, "audit_report.json"), "w"), indent=1)
    print(json.dumps(report["summary"], indent=1))


asyncio.run(main())
